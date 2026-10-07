#!/usr/bin/env python3
"""Rebuild VTK_Field_Report.xlsx (and fill the 'CM render' column of Aum's sweep sheet)
straight from set1.csv / set2.csv, without MATLAB or ParaView.

  python make_field_report.py --set1 set1.csv --set2 set2.csv \
      [--sweep Aums_CM_sweep_2D.csv] [--reference VTK_Field_Report.xlsx] [--out .]

Outputs (in --out):
  renders/set<S>_point<NN>_wv<..>nm_amp<..>pct.png   one thumbnail per point
  VTK_Field_Report_regenerated.xlsx                   same layout as VTK_Field_Report.xlsx
  Aums_CM_sweep_2D_with_renders.xlsx                  only with --sweep (contains SIM data)
--reference compares every render against the ParaView screenshots in the original report.
"""
import argparse
import csv
import io
import os
import sys

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cuberender as cr  # noqa: E402

EMU_PX = 9525
THIN = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="FFD0D1DC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
FONT = Font(name="Arial", size=11)
FONT_B = Font(name="Arial", size=11, bold=True)

# Header characters that were lost ('?') when Aum's CSV was exported.
HEADER_FIX = {"avg ?c? (%)": "avg ⟨c⟩ (%)", "? (nm)": "λ (nm)", "? V1": "Δ V1", "? V3": "Δ V3",
              "? strain (pp)": "Δ strain (pp)", "loop rank ?": "loop rank Δ",
              "SIM cliff ?vf": "SIM cliff Δvf"}


def read_set(path):
    with open(path, newline="") as f:
        return [dict(point=int(r["point"]), avg=round(float(r["avg"]), 6),
                     wv=round(float(r["wv"]), 6), amp=round(float(r["amp"]), 6))
                for r in csv.DictReader(f)]


def render_name(s, r):
    # same pattern as Aum's .vtk names: point01_wv400nm_amp0.56pct (wavelength integer, amp 2 decimals)
    return f"set{s}_point{r['point']:02d}_wv{r['wv']:.0f}nm_amp{r['amp']:.2f}pct.png"


def profile(r):
    # avg and amp are in %, wv in nm, domain 640 nm (render stripe count = 640/wv)
    return cr.sine_profile(r["avg"], r["wv"], r["amp"])


def put_image(ws, path, row, col, px, x_off=0, y_off=0):
    img = XLImage(path)
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=col, row=row, colOff=int(x_off * EMU_PX),
                                                  rowOff=int(y_off * EMU_PX)),
                               ext=XDRPositiveSize2D(px * EMU_PX, px * EMU_PX))
    ws.add_image(img)


def build_report(sets, png, out):
    wb = Workbook()
    wb.remove(wb.active)
    for s, rows in sets.items():
        ws = wb.create_sheet(f"Set {s}")
        for j, (h, w) in enumerate(zip(["Point", "Image", "Avg", "Wavelength", "Amplitude"],
                                       [8.63, 30.63, 10.63, 12.63, 12.63]), 1):
            c = ws.cell(1, j, h)
            c.font, c.fill, c.alignment, c.border = FONT_B, HDR_FILL, CENTER, BORDER
            ws.column_dimensions[get_column_letter(j)].width = w
        for i, r in enumerate(rows, 2):
            ws.row_dimensions[i].height = 110
            vals = [f"pt{r['point']}", None, r["avg"], r["wv"], r["amp"]]
            for j, (v, fmt) in enumerate(zip(vals, ["General", "General", "0.00", "0", "0.00"]), 1):
                c = ws.cell(i, j, v)
                c.font, c.alignment, c.border, c.number_format = FONT, CENTER, BORDER, fmt
            put_image(ws, png[(s, r["point"])], i - 1, 1, 139, 43, 4)
    wb.save(out)


def build_sweep(path, sets, png, out):
    raw = open(path, "rb").read().decode("cp1252", errors="replace")
    table = list(csv.reader(io.StringIO(raw)))
    head = [HEADER_FIX.get(h, h) for h in table[0]]
    body = [r for r in table[1:] if r]
    fmts = {"avg ⟨c⟩ (%)": "0.00", "λ (nm)": "0", "waves / 640 nm": "0.00", "amp (±%)": "0.00",
            "c min (%)": "0.000", "c max (%)": "0.000", "ML strain range": "0.00",
            "SIM strain range": "0.00", "Δ strain (pp)": "0.00", "ML loop area": "0.0",
            "SIM loop area": "0.00", "ML/SIM loop ratio": "0.0", "loop rank Δ": "+0;-0;+0",
            "SIM max vf": "0.000", "SIM cliff Δvf": "0.000", "SIM residual strain": "0.000000"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Set 1"
    ref = {r["point"]: r for r in sets[1]}
    for j, h in enumerate(head, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = FONT_B, HDR_FILL, CENTER, BORDER
        ws.column_dimensions[get_column_letter(j)].width = 24 if h == "CM render" else \
            (62 if h.startswith("CM field") else (28 if h == "SIM notes" else 12))
    ws.row_dimensions[1].height = 42
    for i, row in enumerate(body, 2):
        d = dict(zip(head, row))
        pt = int(d["Point"])
        r = dict(point=pt, avg=float(d["avg ⟨c⟩ (%)"]), wv=float(d["λ (nm)"]),
                 amp=float(d["amp (±%)"]))
        if abs(r["avg"] - ref[pt]["avg"]) > 1e-6 or r["wv"] != ref[pt]["wv"] or \
                abs(r["amp"] - ref[pt]["amp"]) > 1e-6:
            sys.exit(f"sweep row {pt} disagrees with set1.csv: {r} vs {ref[pt]}")
        ws.row_dimensions[i].height = 111.75
        for j, (h, v) in enumerate(zip(head, row), 1):
            try:
                v = float(v) if v.strip() not in ("",) else None
                if v is not None and v == int(v) and h in ("Point", "loop rank Δ"):
                    v = int(v)
            except ValueError:
                pass
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = FONT, CENTER, BORDER
            c.number_format = fmts.get(h, "General")
        put_image(ws, png[(1, pt)], i - 1, 1, 139, 17, 4)
    ws.freeze_panes = "C2"
    wb.save(out)


def compare(reference, png):
    """Mean abs. error (0-255) of every render against the ParaView screenshot."""
    wb = load_workbook(reference)
    errs = []
    for s in (1, 2):
        for im in wb[f"Set {s}"]._images:
            pt = im.anchor._from.row            # 0-based row -> point = row (header is row 0)
            ref = np.array(Image.open(io.BytesIO(im._data())).convert("RGB"), int)
            mine = np.array(Image.open(png[(s, pt)]).convert("RGB"), int)
            errs.append((s, pt, np.abs(ref - mine).mean()))
    for s, pt, e in errs:
        print(f"  set {s} pt {pt:2d}: MAE {e:.2f}/255")
    print(f"  mean MAE {np.mean([e for *_, e in errs]):.2f}/255, worst {max(e for *_, e in errs):.2f}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--set1", required=True)
    ap.add_argument("--set2", required=True)
    ap.add_argument("--sweep")
    ap.add_argument("--reference")
    ap.add_argument("--out", default=".")
    a = ap.parse_args()

    sets = {1: read_set(a.set1), 2: read_set(a.set2)}
    os.makedirs(os.path.join(a.out, "renders"), exist_ok=True)
    png = {}
    for s, rows in sets.items():
        for r in rows:
            p = os.path.join(a.out, "renders", render_name(s, r))
            cr.save_png(profile(r), p)
            png[(s, r["point"])] = p
    print(f"rendered {len(png)} fields")
    build_report(sets, png, os.path.join(a.out, "VTK_Field_Report_regenerated.xlsx"))
    print("wrote VTK_Field_Report_regenerated.xlsx")
    if a.sweep:
        build_sweep(a.sweep, sets, png, os.path.join(a.out, "Aums_CM_sweep_2D_with_renders.xlsx"))
        print("wrote Aums_CM_sweep_2D_with_renders.xlsx")
    if a.reference:
        print("comparison with the ParaView screenshots in", os.path.basename(a.reference))
        compare(a.reference, png)


if __name__ == "__main__":
    main()
